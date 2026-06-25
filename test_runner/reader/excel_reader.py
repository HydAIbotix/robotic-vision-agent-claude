"""
Read kiosk test cases from the Excel spreadsheet into TestCase dicts.

Expected columns (row 1 = headers):
  Local Test ID | Summary | Description | Preconditions | Test Steps |
  Expected Results | Actual Result
"""
from pathlib import Path
import openpyxl
from test_runner.state import TestCase


COLUMN_MAP = {
    "Local Test ID":   "test_id",
    "Summary":         "summary",
    "Description":     "description",
    "Preconditions":   "preconditions",
    "Test Steps":      "steps_raw",
    "Expected Results": "expected_results_raw",
}


def read_test_cases(excel_path: str) -> list[TestCase]:
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    # Build column index from header row
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    col_idx = {
        COLUMN_MAP[h]: i
        for i, h in enumerate(headers)
        if h in COLUMN_MAP
    }

    cases: list[TestCase] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[col_idx.get("test_id", 0)]:
            continue
        cases.append({
            "test_id":              str(row[col_idx["test_id"]]   or ""),
            "summary":              str(row[col_idx["summary"]]   or ""),
            "description":          str(row[col_idx.get("description", 0)] or ""),
            "preconditions":        str(row[col_idx.get("preconditions", 0)] or ""),
            "steps_raw":            str(row[col_idx["steps_raw"]] or ""),
            "expected_results_raw": str(row[col_idx["expected_results_raw"]] or ""),
        })

    print(f"  [READER] Loaded {len(cases)} test cases from {Path(excel_path).name}")
    return cases
